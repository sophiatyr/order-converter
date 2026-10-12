"""Read reference workbooks without changing them; publish layouts for the offline app."""
import sys,pathlib,json,re,struct,datetime
root=pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0,str(root/'scripts/vendor'))
import xlrd,openpyxl
source=root/'CSQ參考資料/訂單作業'
styles=[];style_ids={};templates=[];warnings=[]
def register(style):
 key=json.dumps(style,sort_keys=True,ensure_ascii=False)
 if key not in style_ids:style_ids[key]=len(styles);styles.append(style)
 return style_ids[key]
def rgb(b,index):
 c=b.colour_map.get(index)
 return {'argb':'FF'+''.join(f'{v:02X}' for v in c)} if c else {'argb':'FF000000'}
def xstyle(b,xf):
 f=b.font_list[xf.font_index];a=xf.alignment;bg=xf.background;br=xf.border
 out={'font':{'name':f.name,'size':f.height/20,'bold':bool(f.bold),'italic':bool(f.italic),'strike':bool(f.struck_out),'color':rgb(b,f.colour_index)},'numFmt':b.format_map.get(xf.format_key).format_str if xf.format_key in b.format_map else 'General'}
 if f.underlined:out['font']['underline']='double' if f.underline_type in [2,0x22] else True
 out['alignment']={'horizontal':{1:'left',2:'center',3:'right',4:'fill',5:'justify',6:'centerContinuous',7:'distributed'}.get(a.hor_align,'general'),'vertical':{0:'top',1:'middle',2:'bottom',3:'justify',4:'distributed'}.get(a.vert_align,'bottom'),'wrapText':bool(a.text_wrapped),'shrinkToFit':bool(a.shrink_to_fit),'indent':a.indent_level}
 if a.rotation:out['alignment']['textRotation']='vertical' if a.rotation==255 else a.rotation if a.rotation<=90 else 90-a.rotation
 border={};names={1:'thin',2:'medium',3:'dashed',4:'dotted',5:'thick',6:'double',7:'hair',8:'mediumDashed',9:'dashDot',10:'mediumDashDot',11:'dashDotDot',12:'mediumDashDotDot',13:'slantDashDot'}
 for side in ['left','right','top','bottom']:
  n=getattr(br,side+'_line_style')
  if n:border[side]={'style':names.get(n,'thin'),'color':rgb(b,getattr(br,side+'_colour_index'))}
 if border:out['border']=border
 if bg.fill_pattern:out['fill']={'type':'pattern','pattern':'solid','fgColor':rgb(b,bg.pattern_colour_index),'bgColor':rgb(b,bg.background_colour_index)}
 return out
def xpages(b,index):
 settings={'paperSize':9,'orientation':'portrait','margins':{'left':.7,'right':.7,'top':.75,'bottom':.75,'header':.3,'footer':.3}}
 offset=b._sh_abs_posn[index];data=b.mem
 while offset+4<=len(data):
  code,length=struct.unpack_from('<HH',data,offset);record=data[offset+4:offset+4+length];offset+=4+length
  if code==0x0A:break
  if code in [0x26,0x27,0x28,0x29] and length>=8:settings['margins'][{0x26:'left',0x27:'right',0x28:'top',0x29:'bottom'}[code]]=struct.unpack_from('<d',record)[0]
  if code==0xA1 and length>=34:
   paper,scale,first,fitw,fith,flags,hr,vr,head,foot,copies=struct.unpack_from('<8H2dH',record)
   settings.update({'paperSize':paper or 9,'orientation':'portrait' if flags&2 else 'landscape','scale':scale or 100,'fitToWidth':fitw,'fitToHeight':fith,'firstPageNumber':first or 1});settings['margins'].update(header=head,footer=foot)
 return settings
def xls(file):
 b=xlrd.open_workbook(str(file),formatting_info=True,on_demand=True);result=[]
 for index in range(b.nsheets):
  s=b.sheet_by_index(index)
  meaningful=[(r,c) for r in range(s.nrows) for c in range(s.ncols) if s.cell_value(r,c)!='']
  if not meaningful:continue
  lastrow=max(r for r,c in meaningful)+1;lastcol=max(c for r,c in meaningful)+1;printarea=None
  for n in b.name_obj_list:
   if n.name=='Print_Area' and n.scope==index:
    try:
     _,rl,rh,cl,ch=n.area2d();lastrow=max(lastrow,rh);lastcol=max(lastcol,ch);printarea=f'{xlrd.formula.colname(cl)}{rl+1}:{xlrd.formula.colname(ch-1)}{rh}'
    except:pass
  lastcol=min(lastcol,45);lastrow=min(lastrow,500)
  cells=[]
  for r in range(min(lastrow,s.nrows)):
   for c in range(min(lastcol,s.ncols)):
    cell=s.cell(r,c);style=register(xstyle(b,b.xf_list[cell.xf_index]));value=cell.value
    if cell.ctype==xlrd.XL_CELL_DATE:
     value={'date':xlrd.xldate_as_datetime(value,b.datemode).isoformat()}
    cells.append({'r':r+1,'c':c+1,'v':value if value!='' else None,'s':style})
  page=xpages(b,index);page['printArea']=printarea or f'A1:{xlrd.formula.colname(lastcol-1)}{lastrow}'
  result.append({'name':s.name,'rowCount':lastrow,'colCount':lastcol,'cells':cells,'merges':[[rl+1,cl+1,rh,ch] for rl,rh,cl,ch in s.merged_cells if rh<=lastrow and ch<=lastcol],'columns':[{'width':s.colinfo_map[c].width/256 if c in s.colinfo_map else (s.defcolwidth or 8),'hidden':bool(s.colinfo_map[c].hidden) if c in s.colinfo_map else False} for c in range(lastcol)],'rows':[{'height':s.rowinfo_map[r].height/20 if r in s.rowinfo_map else s.default_row_height/20,'hidden':bool(s.rowinfo_map[r].hidden) if r in s.rowinfo_map else False} for r in range(lastrow)],'pageSetup':page,'showGridLines':bool(s.show_grid_lines)})
 return result
def color(c):
 if c is None:return None
 if c.type=='rgb':return {'argb':str(c.rgb)}
 if c.type=='indexed':
  from openpyxl.styles.colors import COLOR_INDEX
  return {'argb':COLOR_INDEX[c.indexed]} if c.indexed<len(COLOR_INDEX) else {'argb':'FF000000'}
 if c.type=='theme':return {'theme':c.theme,'tint':c.tint}
 return None
def xlsx(file):
 b=openpyxl.load_workbook(file,data_only=False);cached=openpyxl.load_workbook(file,data_only=True);result=[]
 for s in b:
  meaningful=[c for row in s for c in row if c.value is not None]
  if not meaningful:continue
  lr=max(c.row for c in meaningful);lc=max(c.column for c in meaningful);cells=[]
  for row in s.iter_rows(max_row=lr,max_col=lc):
   for c in row:
    f=c.font;a=c.alignment;style={'font':{'name':f.name or 'Arial','size':f.sz or 11,'bold':bool(f.b),'italic':bool(f.i),'color':color(f.color) or {'argb':'FF000000'}},'numFmt':c.number_format,'alignment':{'horizontal':a.horizontal or 'general','vertical':'middle' if a.vertical=='center' else a.vertical or 'bottom','wrapText':bool(a.wrap_text),'shrinkToFit':bool(a.shrink_to_fit)}}
    if f.u:style['font']['underline']=f.u
    border={}
    for side in ['left','right','top','bottom']:
     bs=getattr(c.border,side)
     if bs and bs.style:border[side]={'style':bs.style,'color':color(bs.color) or {'argb':'FF000000'}}
    if border:style['border']=border
    if c.fill.patternType:style['fill']={'type':'pattern','pattern':c.fill.patternType,'fgColor':color(c.fill.fgColor) or {'argb':'FFFFFFFF'},'bgColor':color(c.fill.bgColor) or {'argb':'FFFFFFFF'}}
    v=c.value
    if isinstance(v,datetime.datetime):v={'date':v.isoformat()}
    elif c.data_type=='f':v={'formula':v[1:],'result':cached[s.title].cell(c.row,c.column).value or 0}
    cells.append({'r':c.row,'c':c.column,'v':v,'s':register(style)})
  result.append({'name':s.title,'rowCount':lr,'colCount':lc,'cells':cells,'merges':[[m.min_row,m.min_col,m.max_row,m.max_col] for m in s.merged_cells if m.max_col<=lc and m.max_row<=lr],'columns':[{'width':s.column_dimensions[openpyxl.utils.get_column_letter(c)].width,'hidden':s.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden} for c in range(1,lc+1)],'rows':[{'height':s.row_dimensions[r].height or s.sheet_format.defaultRowHeight or 15,'hidden':s.row_dimensions[r].hidden} for r in range(1,lr+1)],'pageSetup':{'paperSize':int(s.page_setup.paperSize or 9),'orientation':s.page_setup.orientation or 'portrait','fitToPage':bool(s.sheet_properties.pageSetUpPr and s.sheet_properties.pageSetUpPr.fitToPage),'fitToWidth':s.page_setup.fitToWidth or 1,'fitToHeight':s.page_setup.fitToHeight or 0,'printArea':str(s.print_area).replace("'"+s.title+"'!",'').replace('$','') if s.print_area else f'A1:{openpyxl.utils.get_column_letter(lc)}{lr}','margins':{k:getattr(s.page_margins,k) for k in ['left','right','top','bottom','header','footer']}},'showGridLines':bool(s.sheet_view.showGridLines)})
 return result
paths=[]
for kind,folder in [('pi','客人PI'),('factory','工廠採購單'),('quantity','訂單資料-數量明細表')]:
 paths.extend((kind,f) for f in (source/folder).rglob('*') if f.suffix.lower() in ['.xls','.xlsx'] and not f.name.startswith('~$') and (kind!='quantity' or 'PO副' in f.name))
paths.append(('overview',source/'CSQ-早見表-2026.xlsx'))
for kind,file in paths:
 try:
  sheets=xls(file) if file.suffix.lower()=='.xls' else xlsx(file)
  for sheet in sheets:
   text=' '.join(str(c['v']) for c in sheet['cells'] if c['v'] is not None)
   itemNos=list(dict.fromkeys(re.findall(r'\b(?:[A-Z]{2,}[A-Z0-9-]*\d[A-Z0-9-]*)\b',text)))
   factoryCodes=re.findall(r'PO[正副]-([A-Z]+)',file.name)
   templates.append({'kind':kind,'file':file.relative_to(source).as_posix(),'scNumbers':list(dict.fromkeys(re.findall(r'CS\d{8}(?:FA|PS|MD|IF|IS)?',text+' '+file.name))),'poNumbers':list(dict.fromkeys(re.findall(r'\d{2}OF\d{6}',text+' '+file.name))),'itemNos':itemNos,'factoryCodes':factoryCodes,'packing':'MULTIPLE PACK' if 'MULTIPLE PACK' in text else 'NORMAL PACK','sheet':sheet})
 except Exception as e:warnings.append({'file':file.name,'reason':type(e).__name__})
out={'version':1,'updatedAt':datetime.datetime.now().isoformat(),'styles':styles,'templates':templates,'warnings':warnings}
(root/'public/csq-document-templates.json').write_text(json.dumps(out,ensure_ascii=False,separators=(',',':')),encoding='utf-8')
print(f'Document layouts: {len(templates)} sheets, {len(styles)} styles, {len(warnings)} unreadable workbooks.')
